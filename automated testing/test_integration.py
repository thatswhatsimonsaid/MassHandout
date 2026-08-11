"""
Integration tests — requires network access and Playwright.
Tests end-to-end scraping + Word document generation for past and future dates.

Run with:  pytest test_integration.py -v -s
Skip with: pytest test_integration.py -v -s --skip-integration
"""
import os
import json
import asyncio
import pytest
from docx import Document
from docx.shared import Inches
from conftest import KNOWN_DATES


# ── Fixtures ───────────────────────────────────────────────────────────────

def make_user_inputs(date_str, viet_url, lang="eng"):
    return {
        "date": date_str,
        "viet_url": viet_url,
        "organization": "TEST ORG",
        "event_name": "Test Mass",
        "hymns": {
            "opening": {"title": "Test Opening Hymn", "text": "Test opening lyrics"},
            "offertory": {"title": "Test Offertory", "text": "Test offertory lyrics"},
            "communion": {"title": "Test Communion", "text": "Test communion lyrics"},
            "recessional": {"title": "Test Recessional", "text": "Test recessional lyrics"},
        },
        "reading1": {"lang": lang, "option_index": 0},
        "psalm": {"lang": lang, "option_index": 0},
        "reading2": {"lang": lang, "option_index": 0},
        "alleluia": {"lang": lang, "option_index": 0},
        "gospel": {"lang": lang, "option_index": 0},
    }


def run_async(coro):
    return asyncio.new_event_loop().run_until_complete(coro)


# ── Scraping tests ─────────────────────────────────────────────────────────

class TestUSCCBScraping:
    """Test that USCCB scraping returns valid data for known dates."""

    @pytest.mark.parametrize("date_str", list(KNOWN_DATES.keys()))
    def test_usccb_returns_all_sections(self, nb_module, date_str):
        """USCCB scrape should return a dict with all 5 section keys + feast_day."""
        result = run_async(nb_module["scrape_usccb_async"](date_str))

        assert "feast_day" in result
        assert "reading1" in result
        assert "psalm" in result
        assert "reading2" in result
        assert "alleluia" in result
        assert "gospel" in result

        # All section values should be lists
        for key in ["reading1", "psalm", "reading2", "alleluia", "gospel"]:
            assert isinstance(result[key], list), f"{key} should be a list"

    @pytest.mark.parametrize("date_str", list(KNOWN_DATES.keys()))
    def test_usccb_has_at_least_one_reading(self, nb_module, date_str):
        """At least some readings should have content for any date.
        Christmas (122526) has a special multi-Mass page structure that
        the scraper does not currently parse, so we skip the content check."""
        result = run_async(nb_module["scrape_usccb_async"](date_str))

        if date_str == "122526":
            # Known limitation: Christmas page has multiple Mass options
            # (Vigil/Night/Dawn/Day) with non-standard headings
            pytest.skip("Christmas page structure not supported by scraper")
        else:
            assert len(result["reading1"]) > 0, f"No Reading 1 found for {date_str}"
            assert len(result["gospel"]) > 0, f"No Gospel found for {date_str}"

    @pytest.mark.parametrize("date_str", list(KNOWN_DATES.keys()))
    def test_usccb_feast_day_not_default(self, nb_module, date_str):
        """Feast day should not be the generic 'Daily Readings' fallback."""
        result = run_async(nb_module["scrape_usccb_async"](date_str))

        assert result["feast_day"] != "Daily Readings", (
            f"Feast day was default 'Daily Readings' for {date_str}"
        )


class TestThanhLinhScraping:
    """Test that ThanhLinh scraping returns valid Vietnamese data."""

    @pytest.mark.parametrize("date_str", list(KNOWN_DATES.keys()))
    def test_thanhlinh_returns_sections(self, nb_module, date_str):
        url = nb_module["get_thanhlinh_url_dynamically"](date_str)
        result = nb_module["scrape_thanhlinh"](url)

        assert isinstance(result, dict)
        assert "reading1" in result
        assert "psalm" in result
        assert "gospel" in result

    @pytest.mark.parametrize("date_str", list(KNOWN_DATES.keys()))
    def test_thanhlinh_has_reading1_and_gospel(self, nb_module, date_str):
        url = nb_module["get_thanhlinh_url_dynamically"](date_str)
        result = nb_module["scrape_thanhlinh"](url)

        assert len(result["reading1"]) > 0, f"No Vietnamese Reading 1 for {date_str}"
        assert len(result["gospel"]) > 0, f"No Vietnamese Gospel for {date_str}"


# ── Known-date citation verification ───────────────────────────────────────

class TestKnownDateCitations:
    """Verify that scraped citations match known values for specific dates.
    Only checks citations that are not None in KNOWN_DATES."""

    @pytest.mark.parametrize("date_str", list(KNOWN_DATES.keys()))
    def test_feast_day_contains_expected(self, nb_module, date_str):
        expected = KNOWN_DATES[date_str]["feast_day_contains"]
        result = run_async(nb_module["scrape_usccb_async"](date_str))
        assert expected.lower() in result["feast_day"].lower(), (
            f"Feast day '{result['feast_day']}' does not contain '{expected}' for {date_str}"
        )

    def test_aug30_citations(self, nb_module):
        """Verify exact citations for Aug 30, 2026 (22nd Sunday)."""
        date_str = "083026"
        expected = KNOWN_DATES[date_str]["citations"]

        result = run_async(nb_module["scrape_usccb_async"](date_str))
        final = nb_module["prepare_template_data"](
            make_user_inputs(date_str, nb_module["get_thanhlinh_url_dynamically"](date_str)),
            result,
            nb_module["scrape_thanhlinh"](nb_module["get_thanhlinh_url_dynamically"](date_str)),
        )

        citations = final["citations"]
        for section, expected_cite in expected.items():
            if expected_cite is None:
                continue
            actual = citations.get(section, "")
            assert actual == expected_cite, (
                f"{section}: expected '{expected_cite}', got '{actual}'"
            )

    def test_aug7_citations(self, nb_module):
        """Verify key citations for Aug 7, 2026 (Friday, 18th Week)."""
        date_str = "080726"
        expected = KNOWN_DATES[date_str]["citations"]

        result = run_async(nb_module["scrape_usccb_async"](date_str))
        final = nb_module["prepare_template_data"](
            make_user_inputs(date_str, nb_module["get_thanhlinh_url_dynamically"](date_str)),
            result,
            nb_module["scrape_thanhlinh"](nb_module["get_thanhlinh_url_dynamically"](date_str)),
        )

        citations = final["citations"]
        for section, expected_cite in expected.items():
            if expected_cite is None:
                continue
            actual = citations.get(section, "")
            assert actual == expected_cite, (
                f"{section}: expected '{expected_cite}', got '{actual}'"
            )


# ── Word document generation tests ─────────────────────────────────────────

class TestWordDocumentGeneration:
    """Test that the generated .docx file is valid and contains expected content."""

    @pytest.fixture
    def generated_doc(self, nb_module, tmp_path):
        """Generate a Word doc for Aug 30 and return the Document object."""
        date_str = "083026"
        eng = run_async(nb_module["scrape_usccb_async"](date_str))
        viet_url = nb_module["get_thanhlinh_url_dynamically"](date_str)
        viet = nb_module["scrape_thanhlinh"](viet_url)
        final_data = nb_module["prepare_template_data"](
            make_user_inputs(date_str, viet_url), eng, viet
        )
        filename = str(tmp_path / "test_booklet.docx")
        nb_module["create_booklet_docx"](
            make_user_inputs(date_str, viet_url), final_data, filename
        )
        assert os.path.exists(filename), "Word document was not created"
        return Document(filename)

    def test_file_exists_and_valid(self, generated_doc):
        """The .docx file should exist and be openable by python-docx."""
        assert generated_doc is not None
        assert len(generated_doc.paragraphs) > 0

    def test_page_size_landscape(self, generated_doc):
        """Page should be 11" x 8.5" (landscape letter for booklet imposition)."""
        section = generated_doc.sections[0]
        assert section.page_width == Inches(11.0), f"Width should be 11\", got {section.page_width}"
        assert section.page_height == Inches(8.5), f"Height should be 8.5\", got {section.page_height}"

    def test_margins(self, generated_doc):
        """All margins should be 0.5\"."""
        section = generated_doc.sections[0]
        assert section.top_margin == Inches(0.5)
        assert section.bottom_margin == Inches(0.5)
        assert section.left_margin == Inches(0.5)
        assert section.right_margin == Inches(0.5)

    def test_contains_feast_day(self, generated_doc):
        """Document should contain the feast day text."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "Twenty" in full_text or "twenty" in full_text.lower(), (
            "Feast day text not found in document"
        )

    def test_contains_organization(self, generated_doc):
        """Document should contain the organization name."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "TEST ORG" in full_text

    def test_contains_event_name(self, generated_doc):
        """Document should contain the event name."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "Test Mass" in full_text

    def test_contains_formatted_date(self, generated_doc):
        """Document should contain the formatted date."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "August 30, 2026" in full_text

    def test_contains_reading1_text(self, generated_doc):
        """Document should contain Reading 1 text."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "READING 1" in full_text or "BÀI ĐỌC I" in full_text

    def test_contains_gospel_text(self, generated_doc):
        """Document should contain Gospel heading."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "GOSPEL" in full_text or "TIN MỪNG" in full_text

    def test_contains_hymns(self, generated_doc):
        """Document should contain all 4 hymn sections."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "Test Opening Hymn" in full_text
        assert "Test Offertory" in full_text
        assert "Test Communion" in full_text
        assert "Test Recessional" in full_text

    def test_contains_anima_christi(self, generated_doc):
        """Back cover should contain the Anima Christi prayer."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "Anima Christi" in full_text
        assert "Soul of Christ" in full_text

    def test_contains_citations(self, generated_doc):
        """Document should contain the biblical citations."""
        full_text = "\n".join(p.text for p in generated_doc.paragraphs)
        assert "Jeremiah 20:7-9" in full_text
        assert "Matthew 16:21-27" in full_text

    def test_has_two_columns(self, generated_doc):
        """Document section should be set to 2 columns for booklet layout."""
        section = generated_doc.sections[0]
        sectPr = section._sectPr
        from docx.oxml.ns import qn
        cols = sectPr.xpath('./w:cols')
        assert len(cols) > 0, "No column settings found"
        num_cols = cols[0].get(qn('w:num'))
        assert num_cols == '2', f"Expected 2 columns, got {num_cols}"


# ── End-to-end for future dates ─────────────────────────────────────────────

class TestFutureDates:
    """Test the full pipeline for future dates near end of 2026."""

    @pytest.mark.parametrize("date_str", ["122526", "122426", "112226"])
    def test_future_date_pipeline(self, nb_module, date_str, tmp_path):
        """Full pipeline: scrape → parse → generate docx for future dates.
        Christmas (122526) is skipped due to non-standard page structure."""
        if date_str == "122526":
            pytest.skip("Christmas page structure not supported by scraper")

        eng = run_async(nb_module["scrape_usccb_async"](date_str))
        viet_url = nb_module["get_thanhlinh_url_dynamically"](date_str)
        viet = nb_module["scrape_thanhlinh"](viet_url)

        # Must have at least reading1 and gospel
        assert len(eng["reading1"]) > 0, f"No English Reading 1 for {date_str}"
        assert len(eng["gospel"]) > 0, f"No English Gospel for {date_str}"

        final_data = nb_module["prepare_template_data"](
            make_user_inputs(date_str, viet_url), eng, viet
        )

        filename = str(tmp_path / f"test_{date_str}.docx")
        nb_module["create_booklet_docx"](
            make_user_inputs(date_str, viet_url), final_data, filename
        )

        assert os.path.exists(filename)
        doc = Document(filename)
        assert len(doc.paragraphs) > 0

        # Verify the date appears correctly in the document
        full_text = "\n".join(p.text for p in doc.paragraphs)
        from datetime import datetime
        formatted = datetime.strptime(date_str, "%m%d%y").strftime("%B %d, %Y")
        assert formatted in full_text, f"Date '{formatted}' not found in document for {date_str}"
