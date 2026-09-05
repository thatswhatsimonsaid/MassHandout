### Packages ###
from datetime import datetime
import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

### SET COLUMN SPACING ###
def set_col_spacing(section, num_cols=2, space_in_inches=1.2):
    """Configures the number of columns and gutter spacing for a given document section."""
    sectPr = section._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), str(num_cols))
        cols[0].set(qn('w:space'), str(int(space_in_inches * 1440)))


### ADD HEADING WITH CITATION ###
def add_heading_with_citation(doc, label, citation):
    """Creates a heading line with the section label on the left and a right-aligned citation."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    p.paragraph_format.tab_stops.add_tab_stop(Inches(4.4), WD_TAB_ALIGNMENT.RIGHT)
    
    r_label = p.add_run(label)
    r_label.font.bold = True
    r_label.font.size = Pt(10.5)
    r_label.font.color.rgb = RGBColor(0, 0, 0)
    
    if citation:
        r_tab = p.add_run("\t")
        r_cite = p.add_run(citation)
        r_cite.font.bold = False
        r_cite.font.italic = False
        r_cite.font.size = Pt(9.5)
        r_cite.font.color.rgb = RGBColor(80, 80, 80)


### ADD READING SECTION ###
def add_reading_section(doc, label, citation, text):
    """Renders a standard scripture reading section with a right-aligned citation and body text."""
    if not text:
        return
        
    add_heading_with_citation(doc, label, citation)
        
    p_txt = doc.add_paragraph()
    p_txt.paragraph_format.space_after = Pt(3)
    r_txt = p_txt.add_run(text)
    r_txt.font.size = Pt(9.5)


### ADD PSALM SECTION ###
def add_psalm_section(doc, final_data, lang_choice):
    """Renders the responsorial psalm section including response text and formatted verses with response indicators."""
    citations = final_data.get("citations", {})
    psalm_data = final_data["viet"].get("psalm") if lang_choice == "viet" else final_data["eng"].get("psalm")
    
    if not psalm_data:
        return
        
    add_heading_with_citation(doc, "RESPONSORIAL PSALM / ĐÁP CA", citations.get("psalm", ""))
    
    response_text = psalm_data.get("response", "")
    if response_text:
        p_resp = doc.add_paragraph()
        p_resp.paragraph_format.space_after = Pt(4)
        r_run = p_resp.add_run(response_text)
        r_run.font.bold = True
        r_run.font.size = Pt(9.5)
        
    verses_dict = psalm_data.get("verses", {})
    for v_key, verse_text in verses_dict.items():
        p_v = doc.add_paragraph()
        p_v.paragraph_format.space_after = Pt(4)
        p_v.paragraph_format.space_before = Pt(0)
        p_v.paragraph_format.left_indent = Inches(0.15)
        
        v_run = p_v.add_run(verse_text + " ")
        v_run.font.size = Pt(9.5)
        
        r_run = p_v.add_run("R.")
        r_run.font.bold = True
        r_run.font.size = Pt(9.5)


### ADD HYMN SECTION ###
def add_hymn_section(doc, label, hymn_data):
    """Renders a liturgical hymn section using either structured chorus/verses or vanilla text format."""
    if not hymn_data or not hymn_data.get("title"):
        return
        
    add_heading_with_citation(doc, label.upper(), "")
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run(f"Thánh Ca: {hymn_data.get('title')}")
    r_title.font.bold = True
    r_title.font.size = Pt(9.5)
    
    if "text" in hymn_data:
        p_txt = doc.add_paragraph()
        p_txt.paragraph_format.space_after = Pt(4)
        r_txt = p_txt.add_run(hymn_data.get("text"))
        r_txt.font.size = Pt(9.5)
    else:
        chorus = hymn_data.get("chorus", "")
        if chorus:
            p_ch = doc.add_paragraph()
            p_ch.paragraph_format.space_after = Pt(4)
            r_ch = p_ch.add_run(chorus)
            r_ch.font.bold = True
            r_ch.font.size = Pt(9.5)
            
        verses = hymn_data.get("verses", {})
        for v_key, v_text in verses.items():
            p_v = doc.add_paragraph()
            p_v.paragraph_format.space_after = Pt(4)
            p_v.paragraph_format.space_before = Pt(0)
            p_v.paragraph_format.left_indent = Inches(0.15)
            
            r_v = p_v.add_run(v_text)
            r_v.font.size = Pt(9.5)
        
    doc.add_paragraph()


### CREATE BOOKLET DOCX ###
def create_booklet_docx(user_inputs, final_data, filename="Mass_Booklet_Imposed.docx"):
    """Orchestrates the creation and imposition of the 4-panel mass booklet Word document across pages and columns."""
    doc = Document()
    
    section = doc.sections[0]
    section.page_width = Inches(11.0)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    set_col_spacing(section, num_cols=2, space_in_inches=1.2)
    
    hymns = final_data.get("hymns", {})
    citations = final_data.get("citations", {})
    
    # --- PANEL 4: BACK COVER (Left side of Page 1) ---
    p_back_title = doc.add_paragraph()
    r_bt = p_back_title.add_run("CÂU NGUYỆN & KẾT LỄ\n")
    r_bt.font.bold = True
    r_bt.font.size = Pt(11)
    
    p_body = doc.add_paragraph()
    p_body.paragraph_format.space_after = Pt(3)
    r_body = p_body.add_run(
        "Anima Christi:\n"
        "Soul of Christ, sanctify me.\n"
        "Body of Christ, save me.\n"
        "Blood of Christ, embolden me.\n"
        "Water from the side of Christ, wash me.\n"
        "Passion of Christ, strengthen me.\n"
        "O good Jesus, hear me.\n"
        "Within your wounds hide me.\n"
        "Never permit me to be parted from you.\n"
        "From the evil Enemy defend me.\n"
        "At the hour of my death call me and bid me come to you,\n"
        "that with your Saints I may praise you for age upon age.\n"
        "Amen.\n"
    )
    r_body.font.size = Pt(9.5)
    
    if "communion" in hymns:
        add_hymn_section(doc, "Communion Hymn / Thánh Ca Hiệp Lễ", hymns["communion"])
        
    if "recessional" in hymns:
        add_hymn_section(doc, "Recessional Hymn / Thánh Ca Kết Lễ", hymns["recessional"])
        
    doc.add_paragraph()
    
    p_break1 = doc.add_paragraph()
    p_break1.add_run().add_break(docx.enum.text.WD_BREAK.COLUMN)
    
    # --- PANEL 1: FRONT COVER (Right side of Page 1) ---
    p_front = doc.add_paragraph()
    p_front.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    org_text = user_inputs.get("organization", "ORGANIZATION")
    r_org = p_front.add_run(f"{org_text}\n")
    r_org.font.size = Pt(12)
    r_org.font.bold = True
    r_org.font.color.rgb = RGBColor(100, 100, 100)
    
    event_text = user_inputs.get("event_name", "Mass Booklet")
    r_event = p_front.add_run(f"{event_text}\n")
    r_event.font.size = Pt(14)
    r_event.font.bold = True
    
    r_feast = p_front.add_run(f"{final_data.get('feast_day', 'Mass Readings')}\n")
    r_feast.font.size = Pt(11.5)
    
    formatted_date = datetime.strptime(user_inputs.get("date"), "%m%d%y").strftime("%B %d, %Y")
    r_date = p_front.add_run(f"{formatted_date}\n\n")
    r_date.font.size = Pt(10)
    r_date.font.italic = True
    
    if "opening" in hymns:
        add_hymn_section(doc, "Opening Hymn / Thánh Ca Nhập Lễ", hymns["opening"])
    
    p_break2 = doc.add_paragraph()
    p_break2.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    
    # --- PANEL 2: INSIDE LEFT (Page 2 of booklet) ---
    sections_left = [
        ("reading1", "READING 1 / BÀI ĐỌC I"),
        ("psalm", "RESPONSORIAL PSALM / ĐÁP CA"),
        ("reading2", "READING 2 / BÀI ĐỌC II")
    ]
    
    for key, label in sections_left:
        conf = user_inputs.get(key, {"lang": "eng"})
        lang = conf.get("lang", "eng")
        
        if key == "psalm":
            add_psalm_section(doc, final_data, lang)
        else:
            text_body = final_data["viet"].get(key) if lang == "viet" else final_data["eng"].get(key)
            if text_body:
                add_reading_section(doc, label, citations.get(key, ""), text_body)
                doc.add_paragraph()
            
    p_break3 = doc.add_paragraph()
    p_break3.add_run().add_break(docx.enum.text.WD_BREAK.COLUMN)
    
    # --- PANEL 3: INSIDE RIGHT (Page 3 of booklet) ---
    sections_right = [
        ("alleluia", "ALLELUIA"),
        ("gospel", "GOSPEL / TIN MỪNG")
    ]
    
    for key, label in sections_right:
        conf = user_inputs.get(key, {"lang": "eng"})
        lang = conf.get("lang", "eng")
        text_body = final_data["viet"].get(key) if lang == "viet" else final_data["eng"].get(key)
        if text_body:
            display_label = label
            if key == "alleluia":
                eng_all_text = final_data["eng"].get("alleluia", "") or ""
                if "alleluia" not in eng_all_text.lower():
                    display_label = "GOSPEL ACCLAMATION / TUNG HÔ TIN MỪNG"
                    
            add_reading_section(doc, display_label, citations.get(key, ""), text_body)
            doc.add_paragraph()
            
    if "offertory" in hymns:
        add_hymn_section(doc, "Offertory Hymn / Thánh Ca Tiến Lễ", hymns["offertory"])

    doc.save(filename)
    print(f"Booklet handout successfully saved as {filename}!")
